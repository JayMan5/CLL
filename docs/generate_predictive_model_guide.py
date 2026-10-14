"""Generate a plain-language Word guide to the synthetic 30-day model demo."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = ROOT / "COURTLOG_Predictive_Model_Explained.docx"


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def add_callout(document: Document, heading: str, body: str, *, fill: str = "FFF2CC") -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    title = paragraph.add_run(heading + "\n")
    title.bold = True
    title.font.color.rgb = RGBColor(125, 76, 0)
    paragraph.add_run(body)
    document.add_paragraph()


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_document() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Aptos Display"

    title = document.add_heading("CourtLOG Predictive Model", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("A plain-language guide to the synthetic 30-day demo")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    date_line = document.add_paragraph("For the CS team  |  14 October 2026")
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_callout(
        document,
        "Important: demonstration only",
        "This model was trained and tested only on generated fictional data. Its score is not a reliable prediction for a real court case, not a legal deadline or rule, and must not be used to make decisions about hearings, custody, enforcement, or a judge's work.",
    )

    document.add_heading("1. What question does the demo ask?", level=1)
    document.add_paragraph(
        "At a particular snapshot in time, the model estimates whether the record will show no qualifying progress during the following 30 days. In this demo, “progress” means at least one of these recorded events:"
    )
    add_bullet(document, "A hearing-log entry with the outcome Heard or Adjourned.")
    add_bullet(document, "A judgment delivery record.")
    document.add_paragraph(
        "A QR scan, file movement, or internal/administrative change does not count as progress for this model. This 30-day definition was chosen for the software demonstration. It is not a legal time limit and has not been approved as a real-world court performance measure."
    )

    document.add_heading("2. What information does it look at?", level=1)
    inputs = document.add_table(rows=1, cols=2)
    inputs.style = "Light Shading Accent 1"
    set_cell_text(inputs.rows[0].cells[0], "Information", bold=True)
    set_cell_text(inputs.rows[0].cells[1], "Plain-language meaning", bold=True)
    rows = [
        ("Case type and court", "The fictional category and fictional court name in the generated demonstration records."),
        ("Adjournments", "The recorded case-level adjournment count at the snapshot."),
        ("Days since filing", "How long ago the case was filed, measured at the snapshot."),
        ("Days since last recorded hearing or judgment", "The model uses the latest Heard/Adjourned hearing-log date or judgment-delivery date. QR scans and internal/admin activity are ignored. If no qualifying event is recorded, filing age is used as a fallback."),
    ]
    for label, description in rows:
        cells = inputs.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], description)

    document.add_heading("3. How does it produce a score?", level=1)
    for text in [
        "A computer script creates fictional cases, snapshots, feature values, and future waiting times. It does not read CourtLOG's repository case-data files or a court database.",
        "The script labels each fictional case: 1 means the simulation has no qualifying progress in the next 30 days; 0 means it does. These are simulated labels, not observed court outcomes.",
        "A Logistic Regression model looks for patterns in those generated examples and combines them into a score. In simple terms, it learns weights for the input details that help distinguish the two simulated labels.",
        "The score is the model's estimate for that fictional label. A displayed 70% is not a 70% chance that a real case will be delayed. The prototype's 70% flag threshold is retained from earlier code and has not been validated.",
    ]:
        add_bullet(document, text)

    document.add_heading("4. What has been tested?", level=1)
    document.add_paragraph(
        "The reproducible run made 1,500 fictional case snapshots using seed 20261014. The evaluation kept later-starting fictional cases out of training: 1,199 rows were used to train the evaluation model and 301 later-cohort rows were held out for testing. The holdout began at 4 December 2025. This checks that the training and testing code works on a forward-time split; it does not recreate or validate real court practice."
    )
    metrics = document.add_table(rows=1, cols=3)
    metrics.style = "Light Shading Accent 1"
    for cell, heading in zip(metrics.rows[0].cells, ("Check", "Demo result", "What it means here")):
        set_cell_text(cell, heading, bold=True)
    metric_rows = [
        ("ROC AUC", "0.745", "On these invented cases, the model ranked simulated no-progress cases above progress cases better than chance."),
        ("Average precision", "0.736", "The model separated positive labels in this synthetic set; the prior baseline was 0.495."),
        ("Brier score", "0.207", "A probability-error check; lower is better. The synthetic prior baseline was 0.250."),
    ]
    for result in metric_rows:
        cells = metrics.add_row().cells
        for cell, value in zip(cells, result):
            set_cell_text(cell, value)
    document.add_paragraph(
        "These numbers describe only the invented simulation. They must not be quoted as accuracy or reliability on real court cases. The generator deliberately creates patterns for the model to learn, so good synthetic results do not prove that those patterns exist in real courts."
    )

    document.add_heading("5. Which account can test it?", level=1)
    account = document.add_table(rows=0, cols=2)
    account.style = "Light Shading Accent 1"
    for label, value in [
        ("Role", "Deputy Chief Registrar"),
        ("Username", "dcr"),
        ("Demo password", "123"),
        ("Demo scope", "FHC Abuja Court 4 — all divisions"),
    ]:
        cells = account.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], value)
    document.add_paragraph(
        "These are intentionally weak demo credentials. They work only in a fresh, disposable DEMO_MODE=true database. Never use them on a public, staging, or production system. The Deputy Chief Registrar role is used because the prediction screen and refresh endpoint are restricted to that role."
    )

    document.add_heading("6. How to check that the right model is running", level=1)
    for text in [
        "Use the AI Risk Intelligence screen and choose Recalculate Prototype Scores.",
        "Open Why? for a case. For the intended synthetic artifact, the source should say “Synthetic 30-day no-progress demo model” and the explanation should show the 30-day target and fictional-data warning.",
        "If the source says “Deterministic heuristic fallback,” the trained synthetic artifact was not loaded. The heuristic is not a trained 30-day model. Confirm that the artifact was generated and that the disposable local app has DEMO_MODE=true and COURTLOG_MODEL_DIR=.local/courtlog-synthetic-demo/model set before restarting it.",
    ]:
        add_bullet(document, text)
    document.add_paragraph(
        "The generated CSV, metadata, and model file are kept in the ignored .local/courtlog-synthetic-demo/ directory and are not committed. README.md contains the repeatable generation and training commands."
    )

    document.add_heading("7. What is still unknown?", level=1)
    add_bullet(document, "Whether the same relationships hold for actual court cases; this has not been tested.")
    add_bullet(document, "Whether the score is well-calibrated, fair across case groups, or useful to staff.")
    add_bullet(document, "Whether 30 days or the 70% flag threshold is appropriate for any real court workflow.")
    add_bullet(document, "Whether any real case data may be used; permission has not been established for the tracked source files.")
    add_callout(
        document,
        "Bottom line",
        "The demo proves the software can generate fictional data, train a model, evaluate it on later fictional cases, and report the score source. It does not prove that the model can predict real case progress.",
        fill="E2F0D9",
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CourtLOG synthetic model demonstration • Not for real-case decisions")

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(f"Generated {build_document()}")
