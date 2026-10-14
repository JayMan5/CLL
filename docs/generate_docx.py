"""Build the current, evidence-based CourtLOG overview document.

Run with the repository development environment after installing requirements-dev.txt.
The generated document is docs/CourtLOG_Overview.docx.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = ROOT / "CourtLOG_Overview.docx"


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def build_document() -> Path:
    document = Document()
    title = document.add_heading("CourtLOG — Competition Prototype Overview", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "CourtLOG is a competition-team software prototype, not a government entity, "
        "certified public-sector service, or production-ready court system. Feature "
        "development is authorized and does not wait on government approval. Use "
        "fictional/demo records unless written permission authorizes specific real data."
    )

    document.add_heading("1. What the prototype contains", level=1)
    add_bullet(
        document,
        "A FastAPI backend with authenticated, role-scoped case workflows, a persistent "
        "SQLite option, and audit events. Server authorization remains authoritative.",
    )
    add_bullet(
        document,
        "A responsive browser interface and installable Sheriff QR-capture PWA. The PWA "
        "does not queue custody scans offline; an offline scan is not a confirmed check-in.",
    )
    add_bullet(
        document,
        "The approved first-pilot role model covers Chief Judge, Chief Registrar, Deputy Chief Registrar, "
        "Registrar, Filing Clerk, Registry Clerk, Court Clerk, Bailiff, and Sheriff; optional Presiding Judge "
        "access is read-only. Court Clerk requires a server-side organization/deployment entitlement, which "
        "does not process or verify payment. Legacy Clerk accounts are blocked from sign-in until a Chief "
        "Registrar explicitly remaps them; Cashier/Fees Clerk, Counsel, and Counsel's Clerk remain unavailable "
        "until separately approved. See docs/COURTLOG_ROLE_ACCESS_MATRIX.md.",
    )
    add_bullet(
        document,
        "Experimental delay-risk scores have a clearly labelled synthetic demo option. "
        "It uses only generated, fictional case snapshots and asks whether no Heard/Adjourned "
        "hearing event or judgment delivery is recorded in the next 30 days; QR scans and "
        "internal/admin changes do not count. This is a software-demo definition, not a legal "
        "deadline or a court-derived outcome. Its labels are simulated, and the local model file "
        "is kept outside the repository and loads only in a disposable demo. When no model is "
        "loaded, the app uses a separately labelled heuristic fallback. The app measures time "
        "since the last recorded hearing or judgment, but excludes custody scans and admin activity. "
        "The legacy tracked CSV has rule-generated labels calibrated from real-derived distributions "
        "with unverified permission and is never selected by default or treated as synthetic. "
        "No model has been trained on observed court outcomes, no real court data was used, and there "
        "is no independent validation against real cases. Scores are not legal findings and must not "
        "drive adjudicative or other consequential decisions. See docs/ML_EVALUATION_READINESS.md.",
    )
    add_bullet(
        document,
        "Workflow timers, execution prompts, and document output are configured "
        "prototypes. Do not describe them as legally approved rules or forms; legal-rule "
        "changes require Law Lead sign-off.",
    )

    document.add_heading("2. WhatsApp delivery: current state", level=1)
    document.add_paragraph(
        "The code contains an opt-in-gated Meta WhatsApp Cloud API adapter for one-to-one "
        "generic template notices. The implemented adjournment event is the notification "
        "trigger; the app does not send sheriff overdue-writ alerts or join existing "
        "registry groups. Case numbers, hearing dates, reasons, and case details are not "
        "placed in the provider message payload."
    )
    add_bullet(
        document,
        "A phone number in a case is not consent. The Chief Registrar screen records "
        "existing consent evidence; it does not collect or independently verify consent. "
        "An active opt-in is checked again before sending, and inbound STOP/UNSUBSCRIBE "
        "keywords revoke it.",
    )
    add_bullet(
        document,
        "The adapter is disabled by default. DEMO_MODE=true always blocks real sends. "
        "The signed HTTPS webhook handles provider status callbacks and opt-outs; message "
        "records keep keyed recipient hashes and masked numbers rather than raw numbers.",
    )
    add_bullet(
        document,
        "A provider HTTP acceptance is not proof of delivery. The activity log can show "
        "queued, accepted, sent, delivered, read, failed, simulated, disabled, or unknown. "
        "Unknown network outcomes are not automatically retried because Meta may already "
        "have accepted the message.",
    )

    document.add_heading("3. Controlled Cloud API test setup", level=1)
    document.add_paragraph(
        "A real delivery test requires an external Meta WhatsApp Business setup and a "
        "public HTTPS staging origin. Government approval is not a prerequisite for "
        "team development; Meta account/template setup, applicable Meta terms, and the "
        "recipient's explicit opt-in are prerequisites for a real message."
    )
    for text in [
        "Create/configure the Meta Business account, WhatsApp phone number, Cloud API "
        "app, and an approved generic template with no case-specific variables. Use a "
        "team-controlled test recipient.",
        "Deploy the webhook at https://<staging-host>/api/whatsapp/webhook and subscribe "
        "the Meta app to message/status events. Verify the GET challenge and signed POST "
        "callbacks against the staging deployment.",
        "Provision the variables listed in .env.example through the deployment secret "
        "manager (or an ignored local environment file for local development): "
        "WHATSAPP_ENABLED, WHATSAPP_ACCESS_TOKEN, WHATSAPP_PHONE_NUMBER_ID, "
        "WHATSAPP_API_VERSION, WHATSAPP_TEMPLATE_NAME, WHATSAPP_TEMPLATE_LANGUAGE, "
        "WHATSAPP_APP_SECRET, WHATSAPP_WEBHOOK_VERIFY_TOKEN, and a stable "
        "WHATSAPP_CONSENT_HASH_KEY of at least 32 bytes. Never commit or share secrets.",
        "Keep WHATSAPP_ENABLED=false until the controlled test is ready. For the isolated "
        "live-send staging test only, set DEMO_MODE=false and WHATSAPP_ENABLED=true; "
        "record the recipient's actual opt-in evidence, send the generic test template, "
        "and verify provider callbacks and a STOP opt-out.",
        "Inspect the privacy-minimized activity log. Treat accepted/sent as intermediate "
        "states; confirm delivered/read callbacks before claiming delivery. Investigate "
        "unknown before any manual retry.",
    ]:
        add_bullet(document, text)

    document.add_paragraph(
        "Live Meta delivery has not been verified in this repository environment: no "
        "Meta account, approved template, server secrets, public webhook, or controlled "
        "live recipient was supplied. The automated WhatsApp tests use a mocked provider "
        "and do not send network messages."
    )

    document.add_heading("4. Safe verification and acceptance", level=1)
    add_bullet(
        document,
        "Run .venv/bin/python -m pytest tests/test_security.py -k whatsapp -q to exercise "
        "the mocked send request, opt-in gate, privacy-minimized logging, signed callbacks, "
        "opt-out, concurrency guard, and ambiguous-timeout behavior. This is not a live "
        "delivery test.",
    )
    add_bullet(
        document,
        "For the responsive interface smoke test, install the Playwright Chromium browser "
        "once with `npx playwright install chromium`, then run `npm run test:viewport`. It uses "
        "fictional API fixtures and writes fresh output under the ignored .local/viewport-pass/ "
        "directory for 1600×900, 1366×768, 1024×768, 768×1024, and 390×844. Reference PNGs and "
        "measurements are committed in docs/viewport-evidence/. This headless check "
        "does not replace the actual phone, camera, printer, PWA-install, accessibility, or zoom tests.",
    )
    add_bullet(
        document,
        "Do not train or evaluate predictive behavior with tracked court records or "
        "real-derived distributions without the user's documented permission. Synthetic "
        "tests validate software behavior, not real-world model performance.",
    )
    add_bullet(
        document,
        "Complete phone/tablet, camera, install, printer, HTTPS staging, recovery, "
        "backup/restore, and operational acceptance separately before any pilot claim.",
    )

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(f"Generated {output}")
